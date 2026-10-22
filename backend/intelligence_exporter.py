"""
Modul Export Laporan Intelijen Eksekutif (PDF & Excel) untuk Spy SPSE.
Menghasilkan dossier komprehensif profil kompetitor dan laporan komparasi Head-to-Head.
"""

import io
from datetime import datetime
from typing import Dict, Any, List

# OpenPyXL imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab imports
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch, cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas


class NumberedCanvas(canvas.Canvas):
    """Canvas yang menghitung total halaman secara otomatis (two-pass)."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Header running
        self.drawString(36, 810, "SPY SPSE — DOSSIER INTELIJEN PENGADAAN (CONFIDENTIAL)")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 804, 559, 804)

        # Footer running
        self.line(36, 40, 559, 40)
        now_str = datetime.now().strftime("%d/%m/%Y %H:%M WIB")
        self.drawString(36, 28, f"Dicetak pada: {now_str} | Dokumen Rahasia Khusus Tim Bidding Internal")
        page_text = f"Halaman {self._pageNumber} dari {page_count}"
        self.drawRightString(559, 28, page_text)

        self.restoreState()


def format_rupiah(val: float) -> str:
    """Format float ke format Rupiah standar."""
    if not val:
        return "Rp 0"
    try:
        return f"Rp {val:,.0f}".replace(",", ".")
    except Exception:
        return str(val)


import html
import re

def sanitize_for_pdf(text: Any) -> str:
    """Membersihkan tag HTML mentah (misal: span class badge) dan meng-escape karakter XML untuk ReportLab."""
    if text is None:
        return "-"
    raw_str = str(text)
    # Strip HTML tags
    clean = re.sub(r'<[^>]+>', ' ', raw_str)
    # Unescape HTML entities
    clean = html.unescape(clean)
    # Re-escape only XML special chars for ReportLab
    clean = clean.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    # Collapse multiple whitespaces
    clean = re.sub(r'\s+', ' ', clean)
    return clean.strip() or "-"


# ==============================================================================
# 1. DOSSIER KOMPETITOR TUNGGAL (PDF & EXCEL)
# ==============================================================================

def generate_competitor_dossier_pdf(
    profile: Dict[str, Any],
    pricing: Dict[str, Any],
    client_map: Dict[str, Any],
    partners: Dict[str, Any],
    disqualifications: Dict[str, Any]
) -> bytes:
    """Menghasilkan Dossier PDF Eksekutif Profil Kompetitor."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=48,
        bottomMargin=48
    )

    styles = getSampleStyleSheet()
    
    # Custom Typography Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A')
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#475569')
    )
    h2_style = ParagraphStyle(
        'Heading2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#1E293B')
    )
    body_bold = ParagraphStyle(
        'BodyDarkBold',
        parent=body_style,
        fontName='Helvetica-Bold'
    )
    kpi_title_style = ParagraphStyle(
        'KpiTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B'),
        alignment=1
    )
    kpi_val_style = ParagraphStyle(
        'KpiVal',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0F172A'),
        alignment=1
    )
    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )
    table_cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=table_cell,
        fontName='Helvetica-Bold'
    )
    badge_style = ParagraphStyle(
        'Badge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )

    story = []

    # 1. Header Box
    comp_name = profile.get("name", "Kompetitor Tanpa Nama")
    npwp = profile.get("npwp", "-") or "-"
    address = profile.get("address", "-") or "-"

    story.append(Paragraph("DOSSIER INTELIJEN KOMPETITOR SPSE", subtitle_style))
    story.append(Paragraph(comp_name.upper(), title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"<b>NPWP:</b> {npwp} &nbsp;|&nbsp; <b>Alamat:</b> {address}", subtitle_style))
    story.append(Spacer(1, 10))

    # 2. Executive KPI Cards Table
    total_tenders = profile.get("total_tenders", 0)
    total_wins = profile.get("total_wins", 0)
    win_rate = profile.get("win_rate", 0)
    total_contract = profile.get("total_contract_value", 0)
    char_title = pricing.get("character_title", "Moderat")

    kpi_data = [
        [
            Paragraph("TOTAL TENDER", kpi_title_style),
            Paragraph("TOTAL MENANG", kpi_title_style),
            Paragraph("WIN RATE", kpi_title_style),
            Paragraph("AKUMULASI KONTRAK", kpi_title_style),
            Paragraph("KARAKTER HARGA", kpi_title_style),
        ],
        [
            Paragraph(f"{total_tenders:,}", kpi_val_style),
            Paragraph(f"{total_wins:,}", kpi_val_style),
            Paragraph(f"{win_rate}%", kpi_val_style),
            Paragraph(format_rupiah(total_contract), kpi_val_style),
            Paragraph(f"{char_title}", kpi_val_style),
        ]
    ]

    kpi_table = Table(kpi_data, colWidths=[90, 85, 80, 150, 118])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 12))

    # 3. Profil Karakter Harga & Bidding Sweet Spot
    story.append(Paragraph("1. POLA PENAWARAN HARGA & BIDDING SWEET SPOT", h2_style))
    avg_disc = pricing.get("avg_discount_pct", 0)
    win_disc = pricing.get("win_avg_discount_pct", 0)
    lose_disc = pricing.get("lose_avg_discount_pct", 0)
    char_desc = pricing.get("character_desc", "-")
    raw_sw = pricing.get("sweet_spot")
    if isinstance(raw_sw, dict):
        sweet_spot = f"{raw_sw.get('min', 0)}% – {raw_sw.get('max', 0)}% dari Pagu"
    elif raw_sw:
        sweet_spot = str(raw_sw)
    else:
        sweet_spot = "Menyesuaikan HPS Pokja"

    pricing_text = f"""
    <b>Karakteristik Penawaran:</b> {char_title} — {char_desc}<br/>
    <b>Rata-rata Diskon Keseluruhan:</b> {avg_disc}% dari Pagu<br/>
    <b>Diskon saat Menang:</b> {win_disc}% dari Pagu &nbsp;|&nbsp; <b>Diskon saat Gugur/Kalah:</b> {lose_disc}% dari Pagu<br/>
    <b>Winning Sweet Spot:</b> Kompetitor paling efektif memenangkan tender pada rentang <b>{sweet_spot}</b>.
    """
    story.append(Paragraph(pricing_text, body_style))
    story.append(Spacer(1, 10))

    # 4. Peta Kandang Satker (Client & Satker Mapping)
    story.append(Paragraph("2. PETA KANDANG DINAS & SATUAN KERJA DOMINAN", h2_style))
    satkers = (client_map.get("clients") or client_map.get("satker_list") or [])[:8]
    if satkers:
        satker_table_data = [
            [
                Paragraph("<b>No</b>", table_cell_bold),
                Paragraph("<b>Satuan Kerja / Dinas</b>", table_cell_bold),
                Paragraph("<b>Instansi</b>", table_cell_bold),
                Paragraph("<b>Menang</b>", table_cell_bold),
                Paragraph("<b>Total Kontrak</b>", table_cell_bold),
                Paragraph("<b>Rerata Diskon</b>", table_cell_bold)
            ]
        ]
        for idx, s in enumerate(satkers, 1):
            is_hg = s.get("is_home_ground", False) or s.get("is_homeground", False)
            satker_name = sanitize_for_pdf(s.get("satuan_kerja") or s.get("satker", "-"))
            if is_hg:
                satker_name = f"<b>[KANDANG UTAMA]</b> {satker_name}"
            instansi = sanitize_for_pdf(s.get("instansi", "-"))
            total_wins = s.get("total_wins", 0)
            contract_val = format_rupiah(s.get("total_contract_value", 0))
            avg_disc_val = s.get("avg_discount_pct") if s.get("avg_discount_pct") is not None else s.get("avg_discount", 0)

            satker_table_data.append([
                Paragraph(str(idx), table_cell),
                Paragraph(satker_name, table_cell),
                Paragraph(instansi, table_cell),
                Paragraph(f"{total_wins}x", table_cell),
                Paragraph(contract_val, table_cell),
                Paragraph(f"{avg_disc_val}%", table_cell),
            ])
        
        satker_t = Table(satker_table_data, colWidths=[25, 180, 130, 48, 95, 45], repeatRows=1)
        satker_t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        story.append(satker_t)
    else:
        story.append(Paragraph("<i>Belum ada catatan kemenangan spesifik per dinas/satker yang terekam.</i>", body_style))
    story.append(Spacer(1, 10))

    # 5. Rekan Aliansi & Konsorsium (KSO Network)
    story.append(Paragraph("3. JEJARING REKANAN & ALIANSI (KSO / FREQUENT CO-BIDDERS)", h2_style))
    co_bidders = partners.get("co_bidders", [])[:5]
    if co_bidders:
        cb_data = [
            [
                Paragraph("<b>No</b>", table_cell_bold),
                Paragraph("<b>Nama Rekanan / Mitra</b>", table_cell_bold),
                Paragraph("<b>Frekuensi Tender Bersama</b>", table_cell_bold),
                Paragraph("<b>Indikasi Hubungan</b>", table_cell_bold)
            ]
        ]
        for idx, cb in enumerate(co_bidders, 1):
            cb_data.append([
                Paragraph(str(idx), table_cell),
                Paragraph(sanitize_for_pdf(cb.get("partner_name", "-")), table_cell),
                Paragraph(f"{cb.get('together_count', 0)} kali", table_cell),
                Paragraph(sanitize_for_pdf(cb.get("role_label", "Aliansi Bidding Bersama")), table_cell)
            ])
        cb_t = Table(cb_data, colWidths=[25, 250, 110, 138], repeatRows=1)
        cb_t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        story.append(cb_t)
    else:
        story.append(Paragraph("<i>Tidak terdeteksi riwayat kemitraan tetap atau konsorsium KSO formal.</i>", body_style))
    story.append(Spacer(1, 10))

    # 6. Matriks Kelemahan Pokja & Amunisi Sanggahan
    story.append(Paragraph("4. MATRIKS KELEMAHAN POKJA & AMUNISI SANGGAHAN", h2_style))
    clusters_raw = disqualifications.get("weakness_clusters") or disqualifications.get("weakness_matrix", {}).get("clusters", [])
    if isinstance(clusters_raw, dict):
        clusters_list = list(clusters_raw.values())
    elif isinstance(clusters_raw, list):
        clusters_list = clusters_raw
    else:
        clusters_list = []
    tactical_insights = disqualifications.get("tactical_insights", [])

    if clusters_list:
        matrix_data = [
            [
                Paragraph("<b>Klaster Kelemahan Pokja</b>", table_cell_bold),
                Paragraph("<b>Jumlah Gugur</b>", table_cell_bold),
                Paragraph("<b>Persentase</b>", table_cell_bold),
                Paragraph("<b>Tingkat Risiko</b>", table_cell_bold)
            ]
        ]
        for cl in clusters_list:
            lvl = cl.get("risk_level", "info")
            lvl_color = "#DC2626" if lvl == "danger" else "#D97706" if lvl == "warning" else "#2563EB"
            name = sanitize_for_pdf(cl.get("name", "-"))
            count = cl.get("count", 0)
            pct = cl.get("percentage") if cl.get("percentage") is not None else cl.get("pct", 0)
            matrix_data.append([
                Paragraph(f"<b>{name}</b>", table_cell),
                Paragraph(f"{count}x", table_cell),
                Paragraph(f"{pct}%", table_cell),
                Paragraph(f"<font color='{lvl_color}'><b>{lvl.upper()}</b></font>", table_cell)
            ])
        mat_t = Table(matrix_data, colWidths=[230, 95, 95, 103], repeatRows=1)
        mat_t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        story.append(mat_t)
        story.append(Spacer(1, 6))

        if tactical_insights:
            story.append(Paragraph("<b>Rekomendasi Taktis & Peluang Sanggah Pokja:</b>", body_bold))
            for ins in tactical_insights:
                story.append(Paragraph(f"• {sanitize_for_pdf(ins)}", body_style))
    else:
        story.append(Paragraph("<i>Belum tercatat riwayat evaluasi gugur pada database SPSE.</i>", body_style))
    story.append(Spacer(1, 10))

    # 7. Riwayat Tender Kemenangan Terakhir
    story.append(Paragraph("5. RIWAYAT TENDER KEMENANGAN TERAKHIR", h2_style))
    recents = profile.get("recent_tenders", [])[:10]
    if recents:
        rec_data = [
            [
                Paragraph("<b>No</b>", table_cell_bold),
                Paragraph("<b>Nama Paket Tender</b>", table_cell_bold),
                Paragraph("<b>Instansi / LPSE</b>", table_cell_bold),
                Paragraph("<b>Pagu / Nilai Kontrak</b>", table_cell_bold)
            ]
        ]
        for idx, rec in enumerate(recents, 1):
            kontrak_num = rec.get("harga_kontrak") or 0
            kontrak_val = format_rupiah(kontrak_num) if kontrak_num > 0 else "-"
            pagu_num = rec.get("pagu_val") or rec.get("pagu") or 0
            pagu_str = format_rupiah(pagu_num) if pagu_num > 0 else "-"
            rec_name = sanitize_for_pdf(rec.get("nama_tender", "-"))
            rec_inst = sanitize_for_pdf(rec.get("instansi", "-"))
            rec_data.append([
                Paragraph(str(idx), table_cell),
                Paragraph(rec_name, table_cell),
                Paragraph(rec_inst, table_cell),
                Paragraph(f"Pagu: {pagu_str}<br/><b>Kontrak: {kontrak_val}</b>", table_cell)
            ])
        rec_t = Table(rec_data, colWidths=[25, 230, 130, 138], repeatRows=1)
        rec_t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        story.append(rec_t)

    # Build document with NumberedCanvas
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


def generate_competitor_dossier_excel(
    profile: Dict[str, Any],
    pricing: Dict[str, Any],
    client_map: Dict[str, Any],
    partners: Dict[str, Any],
    disqualifications: Dict[str, Any]
) -> bytes:
    """Menghasilkan Workbook Excel Multi-Sheet Dossier Kompetitor."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    accent_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # --- SHEET 1: RINGKASAN EKSEKUTIF ---
    ws1 = wb.create_sheet(title="Ringkasan Eksekutif")
    ws1.views.sheetView[0].showGridLines = True

    # Title
    ws1.append(["DOSSIER INTELIJEN KOMPETITOR SPSE"])
    ws1.append([profile.get("name", "").upper()])
    ws1.append([f"NPWP: {profile.get('npwp', '-')} | Alamat: {profile.get('address', '-')}"])
    ws1.append([f"Waktu Ekspor: {datetime.now().strftime('%d/%m/%Y %H:%M WIB')}"])
    ws1.append([])

    ws1.cell(row=1, column=1).font = Font(name="Arial", size=11, bold=True, color="64748B")
    ws1.cell(row=2, column=1).font = Font(name="Arial", size=16, bold=True, color="0F172A")
    ws1.cell(row=3, column=1).font = Font(name="Arial", size=9, color="475569")
    ws1.cell(row=4, column=1).font = Font(name="Arial", size=8, italic=True, color="94A3B8")

    # KPI Table
    ws1.append(["METRIK UTAMA", "NILAI"])
    ws1.append(["Total Tender Diikuti", profile.get("total_tenders", 0)])
    ws1.append(["Total Menang", profile.get("total_wins", 0)])
    ws1.append(["Tingkat Kemenangan (Win Rate)", f"{profile.get('win_rate', 0)}%"])
    ws1.append(["Akumulasi Nilai Kontrak (Rp)", profile.get("total_contract_value", 0)])
    ws1.append(["Total Pagu Diikuti (Rp)", profile.get("total_pagu", 0)])
    ws1.append(["Karakter Penawaran Harga", pricing.get("character_title", "-")])
    ws1.append(["Rata-rata Diskon dari Pagu", f"{pricing.get('avg_discount_pct', 0)}%"])
    ws1.append(["Diskon saat Menang", f"{pricing.get('win_avg_discount_pct', 0)}%"])
    raw_sw = pricing.get("sweet_spot")
    if isinstance(raw_sw, dict):
        sweet_spot_str = f"{raw_sw.get('min', 0)}% – {raw_sw.get('max', 0)}% dari Pagu"
    elif raw_sw:
        sweet_spot_str = str(raw_sw)
    else:
        sweet_spot_str = "Menyesuaikan HPS Pokja"

    ws1.append(["Winning Sweet Spot", sweet_spot_str])
    ws1.append([])

    # Styling KPI Table
    ws1.cell(row=6, column=1).font = header_font
    ws1.cell(row=6, column=1).fill = header_fill
    ws1.cell(row=6, column=2).font = header_font
    ws1.cell(row=6, column=2).fill = header_fill

    for r in range(7, 18):
        ws1.cell(row=r, column=1).font = Font(name="Arial", size=9, bold=True)
        ws1.cell(row=r, column=1).border = thin_border
        ws1.cell(row=r, column=2).font = Font(name="Arial", size=9)
        ws1.cell(row=r, column=2).border = thin_border
        if r in [10, 11]:  # Currencies
            ws1.cell(row=r, column=2).number_format = '#,##0'

    # --- SHEET 2: PETA KANDANG SATKER ---
    ws2 = wb.create_sheet(title="Peta Kandang Satker")
    ws2.views.sheetView[0].showGridLines = True
    ws2_headers = ["No", "Satuan Kerja / Dinas", "Instansi", "Total Menang", "Total Nilai Kontrak (Rp)", "Rata-rata Diskon (%)", "Status Kandang"]
    ws2.append(ws2_headers)

    for col_idx in range(1, len(ws2_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    satkers_all = client_map.get("clients") or client_map.get("satker_list") or []
    for idx, s in enumerate(satkers_all, 1):
        is_hg = s.get("is_home_ground", False) or s.get("is_homeground", False)
        status = "KANDANG UTAMA" if is_hg else "Satker Aktif"
        satker_name = s.get("satuan_kerja") or s.get("satker", "-")
        avg_disc = s.get("avg_discount_pct") if s.get("avg_discount_pct") is not None else s.get("avg_discount", 0.0)
        row_data = [
            idx,
            satker_name,
            s.get("instansi", "-"),
            s.get("total_wins", 0),
            s.get("total_contract_value", 0.0),
            avg_disc,
            status
        ]
        ws2.append(row_data)
        r_num = idx + 1
        for c_idx in range(1, len(row_data) + 1):
            c_cell = ws2.cell(row=r_num, column=c_idx)
            c_cell.border = thin_border
            c_cell.font = Font(name="Arial", size=9)
            if c_idx == 5:
                c_cell.number_format = '#,##0'
                c_cell.alignment = Alignment(horizontal="right")
            elif c_idx in [1, 4, 6, 7]:
                c_cell.alignment = Alignment(horizontal="center")

    # --- SHEET 3: MITRA & ALIANSI KSO ---
    ws3 = wb.create_sheet(title="Mitra & Aliansi")
    ws3.views.sheetView[0].showGridLines = True
    ws3_headers = ["No", "Nama Rekanan / Mitra", "Frekuensi Tender Bersama", "Indikasi Peran Hubungan"]
    ws3.append(ws3_headers)

    for col_idx in range(1, len(ws3_headers) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = sub_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, cb in enumerate(partners.get("co_bidders", []), 1):
        row_data = [
            idx,
            cb.get("partner_name", "-"),
            cb.get("together_count", 0),
            cb.get("role_label", "Mitra Co-Bidding")
        ]
        ws3.append(row_data)
        r_num = idx + 1
        for c_idx in range(1, len(row_data) + 1):
            c_cell = ws3.cell(row=r_num, column=c_idx)
            c_cell.border = thin_border
            c_cell.font = Font(name="Arial", size=9)
            if c_idx in [1, 3]:
                c_cell.alignment = Alignment(horizontal="center")

    # --- SHEET 4: MATRIKS KELEMAHAN POKJA ---
    ws4 = wb.create_sheet(title="Kelemahan Pokja")
    ws4.views.sheetView[0].showGridLines = True
    ws4_headers = ["Klaster Kelemahan", "Jumlah Gugur", "Persentase (%)", "Tingkat Risiko"]
    ws4.append(ws4_headers)

    for col_idx in range(1, len(ws4_headers) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = accent_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    clusters_excel_raw = disqualifications.get("weakness_clusters") or disqualifications.get("weakness_matrix", {}).get("clusters", [])
    if isinstance(clusters_excel_raw, dict):
        clusters_excel = list(clusters_excel_raw.values())
    elif isinstance(clusters_excel_raw, list):
        clusters_excel = clusters_excel_raw
    else:
        clusters_excel = []

    row_count = 1
    for cl in clusters_excel:
        row_count += 1
        pct = cl.get("percentage") if cl.get("percentage") is not None else cl.get("pct", 0)
        row_data = [
            cl.get("name", "-"),
            cl.get("count", 0),
            pct,
            cl.get("risk_level", "info").upper()
        ]
        ws4.append(row_data)
        for c_idx in range(1, 5):
            c_cell = ws4.cell(row=row_count, column=c_idx)
            c_cell.border = thin_border
            c_cell.font = Font(name="Arial", size=9)
            if c_idx in [2, 3, 4]:
                c_cell.alignment = Alignment(horizontal="center")

    # --- SHEET 5: TENDER DIMENANGKAN ---
    ws5 = wb.create_sheet(title="Tender Dimenangkan")
    ws5.views.sheetView[0].showGridLines = True
    ws5_headers = ["No", "Nomor Tender", "Nama Paket Tender", "Instansi", "Pagu (Rp)", "Harga Kontrak (Rp)"]
    ws5.append(ws5_headers)

    for col_idx in range(1, len(ws5_headers) + 1):
        cell = ws5.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, rec in enumerate(profile.get("recent_tenders", []), 1):
        pagu_val = rec.get("pagu_val") or rec.get("pagu") or 0
        kontrak_val = rec.get("harga_kontrak") or 0
        row_data = [
            idx,
            rec.get("nomor_pengadaan", "-"),
            rec.get("nama_tender", "-"),
            rec.get("instansi", "-"),
            pagu_val,
            kontrak_val
        ]
        ws5.append(row_data)
        r_num = idx + 1
        for c_idx in range(1, len(row_data) + 1):
            c_cell = ws5.cell(row=r_num, column=c_idx)
            c_cell.border = thin_border
            c_cell.font = Font(name="Arial", size=9)
            if c_idx in [1, 2]:
                c_cell.alignment = Alignment(horizontal="center")
            elif c_idx in [5, 6]:
                c_cell.number_format = '#,##0'
                c_cell.alignment = Alignment(horizontal="right")

    # Auto-fit column widths across all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or '')
                if '\n' in val_str:
                    val_str = max(val_str.split('\n'), key=len)
                max_len = max(max_len, len(val_str))
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()


# ==============================================================================
# 2. LAPORAN KOMPARASI HEAD-TO-HEAD (PDF & EXCEL)
# ==============================================================================

def generate_head_to_head_pdf(data: Dict[str, Any]) -> bytes:
    """Menghasilkan Laporan PDF Eksekutif Komparasi Head-to-Head 2 Perusahaan."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=48,
        bottomMargin=48
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0F172A')
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#64748B')
    )
    h2_style = ParagraphStyle(
        'Heading2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=8,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#1E293B')
    )
    score_big = ParagraphStyle(
        'ScoreBig',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    score_sub = ParagraphStyle(
        'ScoreSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#64748B')
    )
    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )
    table_cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=table_cell,
        fontName='Helvetica-Bold'
    )

    c1 = data.get("comp1", {})
    c2 = data.get("comp2", {})
    h2h = data.get("head_to_head", {})
    meetings = data.get("meetings", [])

    story = []

    # Title Banner
    story.append(Paragraph("LAPORAN KOMPARASI HEAD-TO-HEAD KOMPETITOR", subtitle_style))
    story.append(Paragraph(f"{c1.get('name', 'Pesaing 1')} <font color='#64748B'>VS</font> {c2.get('name', 'Pesaing 2')}", title_style))
    story.append(Spacer(1, 8))

    # Tale of the Tape 3-Column Table
    tot_data = [
        [
            Paragraph(f"<b>{c1.get('name', 'Perusahaan 1')}</b>", table_cell_bold),
            Paragraph("<b>SCORECARD BENTROK</b>", ParagraphStyle('HCenter', parent=table_cell_bold, alignment=1)),
            Paragraph(f"<b>{c2.get('name', 'Perusahaan 2')}</b>", ParagraphStyle('HRight', parent=table_cell_bold, alignment=2))
        ],
        [
            Paragraph(f"Win Rate: <b>{c1.get('win_rate', 0)}%</b> ({c1.get('total_wins', 0)}/{c1.get('total_tenders', 0)})", table_cell),
            Paragraph(f"<b>{h2h.get('comp1_wins', 0)}</b> &nbsp;—&nbsp; <b>{h2h.get('comp2_wins', 0)}</b><br/><font size=7 color='#64748B'>Total {h2h.get('total_meetings', 0)} Pertemuan</font>", ParagraphStyle('ScoreC', parent=table_cell, alignment=1)),
            Paragraph(f"Win Rate: <b>{c2.get('win_rate', 0)}%</b> ({c2.get('total_wins', 0)}/{c2.get('total_tenders', 0)})", ParagraphStyle('RCell', parent=table_cell, alignment=2))
        ],
        [
            Paragraph(f"Kontrak: {format_rupiah(c1.get('total_contract_value', 0))}", table_cell),
            Paragraph(f"Lainnya Menang: {h2h.get('other_wins', 0)}x", ParagraphStyle('ScoreC2', parent=table_cell, alignment=1)),
            Paragraph(f"Kontrak: {format_rupiah(c2.get('total_contract_value', 0))}", ParagraphStyle('RCell2', parent=table_cell, alignment=2))
        ],
        [
            Paragraph(f"Rerata Diskon: <b>{c1.get('avg_discount_pct', 0)}%</b>", table_cell),
            Paragraph(f"Selisih Harga: {h2h.get('avg_discount_diff', 0)}%", ParagraphStyle('ScoreC3', parent=table_cell, alignment=1)),
            Paragraph(f"Rerata Diskon: <b>{c2.get('avg_discount_pct', 0)}%</b>", ParagraphStyle('RCell3', parent=table_cell, alignment=2))
        ],
        [
            Paragraph(f"Karakter: {c1.get('character_label', '-')}", table_cell),
            Paragraph(f"Lebih Murah: {h2h.get('comp1_cheaper_count', 0)} vs {h2h.get('comp2_cheaper_count', 0)}", ParagraphStyle('ScoreC4', parent=table_cell, alignment=1)),
            Paragraph(f"Karakter: {c2.get('character_label', '-')}", ParagraphStyle('RCell4', parent=table_cell, alignment=2))
        ]
    ]

    tot_table = Table(tot_data, colWidths=[195, 133, 195])
    tot_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (1, 1), (1, -1), colors.HexColor('#F1F5F9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(tot_table)
    story.append(Spacer(1, 10))

    # Strategic Verdict Callout
    raw_verdict = h2h.get("verdict_description") or h2h.get("strategic_verdict") or "-"
    raw_title = h2h.get("verdict_title") or ""
    verdict_body = sanitize_for_pdf(raw_verdict)
    verdict_title_clean = sanitize_for_pdf(raw_title) if raw_title else ""
    verdict = f"<b>{verdict_title_clean}</b><br/>{verdict_body}" if verdict_title_clean else verdict_body
    advice = sanitize_for_pdf(h2h.get("advice") or "-")
    verdict_box = [
        [
            Paragraph(f"<b>DYNAMIC AI STRATEGIC VERDICT:</b><br/>{verdict}<br/><br/><b>Rekomendasi Taktis Penawaran:</b> {advice}", body_style)
        ]
    ]
    v_table = Table(verdict_box, colWidths=[523])
    v_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#EFF6FF')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#93C5FD')),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(v_table)
    story.append(Spacer(1, 10))

    # Riwayat Tender Bentrok
    story.append(Paragraph("RIWAYAT PERTEMUAN LELANG (BENTROK LANGSUNG)", h2_style))
    if meetings:
        meet_data = [
            [
                Paragraph("<b>No</b>", table_cell_bold),
                Paragraph("<b>Nama Paket Tender</b>", table_cell_bold),
                Paragraph("<b>Pagu (Rp)</b>", table_cell_bold),
                Paragraph(f"<b>Tawaran {c1.get('name', 'Pesaing 1')[:12]}</b>", table_cell_bold),
                Paragraph(f"<b>Tawaran {c2.get('name', 'Pesaing 2')[:12]}</b>", table_cell_bold),
                Paragraph("<b>Pemenang Riil</b>", table_cell_bold)
            ]
        ]
        for idx, m in enumerate(meetings[:15], 1):
            p1_num = m.get("comp1_price") or 0
            p1_val = format_rupiah(p1_num) if p1_num > 0 else "Gugur / Rp 0"
            p2_num = m.get("comp2_price") or 0
            p2_val = format_rupiah(p2_num) if p2_num > 0 else "Gugur / Rp 0"
            pagu_num = m.get("pagu_val") or m.get("pagu") or 0
            pagu_val = format_rupiah(pagu_num) if pagu_num > 0 else "-"
            m_name = sanitize_for_pdf(m.get("nama_tender", "-"))
            winner = sanitize_for_pdf(m.get("winner_name", "-") or "Belum Ditentukan")
            meet_data.append([
                Paragraph(str(idx), table_cell),
                Paragraph(m_name, table_cell),
                Paragraph(pagu_val, table_cell),
                Paragraph(p1_val, table_cell),
                Paragraph(p2_val, table_cell),
                Paragraph(f"<b>{winner}</b>", table_cell)
            ])
        m_table = Table(meet_data, colWidths=[20, 163, 85, 85, 85, 85], repeatRows=1)
        m_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        story.append(m_table)
    else:
        story.append(Paragraph("<i>Tidak ditemukan lelang di mana kedua perusahaan memasukkan penawaran bersama.</i>", body_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


def generate_head_to_head_excel(data: Dict[str, Any]) -> bytes:
    """Menghasilkan Workbook Excel Laporan Komparasi Head-to-Head."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    c1 = data.get("comp1", {})
    c2 = data.get("comp2", {})
    h2h = data.get("head_to_head", {})
    meetings = data.get("meetings", [])

    # Sheet 1: Scorecard Matchup
    ws1 = wb.create_sheet(title="Scorecard Head-to-Head")
    ws1.views.sheetView[0].showGridLines = True

    ws1.append(["KOMPARASI HEAD-TO-HEAD KOMPETITOR SPSE"])
    ws1.append([f"{c1.get('name', 'Pesaing 1')} vs {c2.get('name', 'Pesaing 2')}"])
    ws1.append([f"Waktu Ekspor: {datetime.now().strftime('%d/%m/%Y %H:%M WIB')}"])
    ws1.append([])

    ws1.cell(row=1, column=1).font = Font(name="Arial", size=11, bold=True, color="64748B")
    ws1.cell(row=2, column=1).font = Font(name="Arial", size=15, bold=True, color="0F172A")
    ws1.cell(row=3, column=1).font = Font(name="Arial", size=8, italic=True, color="94A3B8")

    ws1.append(["PARAMETER KOMPARASI", c1.get("name", "Perusahaan 1"), "MATCHUP SCORE", c2.get("name", "Perusahaan 2")])
    ws1.append(["Skor Kemenangan Bentrok", f"{h2h.get('comp1_wins', 0)} Menang", f"Total {h2h.get('total_meetings', 0)} Bentrok", f"{h2h.get('comp2_wins', 0)} Menang"])
    ws1.append(["Total Tender Diikuti", c1.get("total_tenders", 0), "-", c2.get("total_tenders", 0)])
    ws1.append(["Total Kemenangan SPSE", c1.get("total_wins", 0), "-", c2.get("total_wins", 0)])
    ws1.append(["Tingkat Kemenangan (Win Rate)", f"{c1.get('win_rate', 0)}%", "-", f"{c2.get('win_rate', 0)}%"])
    ws1.append(["Total Nilai Kontrak (Rp)", c1.get("total_contract_value", 0), "-", c2.get("total_contract_value", 0)])
    ws1.append(["Rata-rata Diskon dari Pagu", f"{c1.get('avg_discount_pct', 0)}%", f"Selisih: {h2h.get('avg_discount_diff', 0)}%", f"{c2.get('avg_discount_pct', 0)}%"])
    ws1.append(["Karakter Bidding", c1.get("character_label", "-"), "-", c2.get("character_label", "-")])
    ws1.append(["Frekuensi Penawaran Lebih Murah", f"{h2h.get('comp1_cheaper_count', 0)} kali", "Cheaper Bid", f"{h2h.get('comp2_cheaper_count', 0)} kali"])
    ws1.append(["Dinas / Satker Dominan", c1.get("top_satker", "-"), "-", c2.get("top_satker", "-")])
    ws1.append([])

    raw_verdict_xl = h2h.get("verdict_description") or h2h.get("strategic_verdict") or "-"
    raw_title_xl = h2h.get("verdict_title") or ""
    verdict_str = f"{raw_title_xl}: {raw_verdict_xl}" if raw_title_xl else raw_verdict_xl
    advice_str = h2h.get("advice") or "-"
    ws1.append(["DYNAMIC AI STRATEGIC VERDICT", verdict_str])
    ws1.append(["REKOMENDASI TAKTIS", advice_str])

    # Header styling
    for col_idx in range(1, 5):
        cell = ws1.cell(row=5, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in range(6, 16):
        for c in range(1, 5):
            cell = ws1.cell(row=r, column=c)
            cell.border = thin_border
            cell.font = Font(name="Arial", size=9)
            if c in [1, 3]:
                cell.alignment = Alignment(horizontal="center")
            if r == 10 and c in [2, 4]:
                cell.number_format = '#,##0'

    # Sheet 2: Riwayat Pertemuan Tender
    ws2 = wb.create_sheet(title="Riwayat Lelang Bentrok")
    ws2.views.sheetView[0].showGridLines = True
    ws2_headers = [
        "No", "Nomor Tender", "Nama Paket Tender", "Pagu (Rp)",
        f"Penawaran {c1.get('name', 'Pesaing 1')[:15]}",
        f"Penawaran {c2.get('name', 'Pesaing 2')[:15]}",
        "Pemenang Riil", "Selisih Harga (Rp)", "Selisih (%)", "Pihak Lebih Murah"
    ]
    ws2.append(ws2_headers)

    for col_idx in range(1, len(ws2_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = sub_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, m in enumerate(meetings, 1):
        pagu_num = m.get("pagu_val") or m.get("pagu") or 0
        row_data = [
            idx,
            m.get("nomor_pengadaan", "-"),
            m.get("nama_tender", "-"),
            pagu_num,
            m.get("comp1_price", 0) or 0,
            m.get("comp2_price", 0) or 0,
            m.get("winner_name", "-") or "Belum Ditentukan",
            m.get("price_diff_rp", 0) or 0,
            m.get("price_diff_pct", 0) or 0,
            m.get("cheaper_party", "-")
        ]
        ws2.append(row_data)
        r_num = idx + 1
        for c_idx in range(1, len(row_data) + 1):
            c_cell = ws2.cell(row=r_num, column=c_idx)
            c_cell.border = thin_border
            c_cell.font = Font(name="Arial", size=9)
            if c_idx in [1, 2, 9, 10]:
                c_cell.alignment = Alignment(horizontal="center")
            elif c_idx in [4, 5, 6, 8]:
                c_cell.number_format = '#,##0'
                c_cell.alignment = Alignment(horizontal="right")

    # Auto-fit column widths across all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or '')
                if '\n' in val_str:
                    val_str = max(val_str.split('\n'), key=len)
                max_len = max(max_len, len(val_str))
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()
